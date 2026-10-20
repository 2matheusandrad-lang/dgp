#!/usr/bin/env python3
"""
Importação da "PLANILHA DE CONTROLE DE PESSOAL" (Google Sheets exportado em .xlsx)
para o banco do Sistema de Controle do Efetivo (Supabase/PostgreSQL).

Uso:
    pip install openpyxl
    python importar_planilha.py PLANILHA.xlsx  pasta_saida/

Gera:
    pasta_saida/seed_1_militares.sql   -> seções, militares e inclusões (rodar 1º no SQL Editor)
    pasta_saida/seed_2_detalhes.sql    -> afastamentos, situações, formação, idiomas, punições (rodar 2º)
    pasta_saida/relatorio_importacao.xlsx -> resumo por OME, problemas e duplicidades encontrados

Regras aplicadas (ver README):
  * cada aba de OME é lida pelo NOME do cabeçalho (o 1ºGB tem 41 colunas, as demais 38);
  * aba quebrada (#REF!) é recuperada do bloco correspondente na aba GERAL;
  * a OME de lotação é a aba onde o militar está listado; se a matrícula aparece em
    duas abas, fica na aba cuja coluna "OME" coincide; a outra vai para o relatório;
  * a ordem das linhas em cada aba vira "ordem de antiguidade" (a planilha já é ordenada assim);
  * término do afastamento é recalculado (início + dias - 1); datas impossíveis são descartadas;
  * Físico/Mental só é mantido para tipos de saúde (LTS, RTS, Laudo médico).
"""
import sys, re, uuid, datetime as dt, collections, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ----------------------------------------------------------------------------
ABA_PARA_SIGLA = {
    '1°GB': '1ºGB', '2°GB': '2ºGB', '3°GB': '3ºGB', '4°GB': '4ºGB', '5°GB': '5ºGB',
    '6°GB': '6ºGB', '7°GB': '7ºGB', '9°GB': '9ºGB',
    'ABMG': 'ABMG', 'AECI': 'AECI', 'AJG': 'AjG', 'CAT AGRESTE': 'CAT Agreste',
    'CAT SERTÃO': 'CAT Sertão', 'CAT ZONA DA MATA': 'CAT Zona da Mata', 'CATRMR': 'CAT/RMR',
    'CCS': 'CCS', 'CEAO': 'CEAO', 'CEFD': 'CEFD', 'CI': 'CI', 'CINT': 'CInt', 'CMAN': 'CMan',
    'CPOCPP': 'CPO', 'CPPA': 'CPPA', 'DAS': 'DAS', 'DEIP': 'DEIP', 'DF': 'DF', 'DGO': 'DGO',
    'DGP': 'DGP', 'ADIDO': 'ADIDO', 'DIESP': 'DIEsp', 'DIM': 'DIM', 'DINTER1': 'DInter/1',
    'DINTER2': 'DInter/2', 'DJD': 'DJD', 'DLOG': 'DLog', 'DPLAG': 'DPlaG', 'DTIC': 'DTIC',
    'DVP': 'DVP', 'GBAPH': 'GBAPH', 'GBI': 'GBI', 'GBMAR': 'GBMar', 'GBS': 'GBS',
    'GCG': 'GCG', 'GSG': 'GSG',
}
ABAS_IGNORADAS = {'GERAL', 'ArquitEng Civil', 'CADASTROAVISOS'}

POSTOS = {  # forma normalizada (sem espaços, maiúscula) -> id em posto_graduacao
    'CEL': 1, 'TENCEL': 2, 'MAJ': 3, 'CAP': 4, '1ºTEN': 5, '2ºTEN': 6, 'ASP': 7, 'ASPIRANTE': 7,
    'SUBTEN': 8, '1ºSGT': 9, '2ºSGT': 10, '3ºSGT': 11, 'CB': 12, 'SD': 13,
}
QUADROS = {'QOC', 'QOA', 'QOS', 'QOE', 'QOM'}

TIPOS_AFAST = {
    'FÉRIAS': 'FERIAS', 'FERIAS': 'FERIAS', 'LTS': 'LTS', 'RTS': 'RTS', 'LE': 'LE', 'LTIP': 'LTIP',
    'LTSPF': 'LTSPF', 'LICENÇA MATERNIDADE': 'LIC_MATERNIDADE', 'LICENÇA PATERNIDADE': 'LIC_PATERNIDADE',
    'LICENÇA NÚPCIAS': 'LIC_NUPCIAS', 'DISPENSA COMO RECOMPENSA': 'DISPENSA_RECOMPENSA',
    'DISP. RECOMP.': 'DISPENSA_RECOMPENSA', 'TRÂNSITO': 'TRANSITO', 'TRANSITO': 'TRANSITO',
    'CURSO': 'CURSO', 'CURSO OPERACIONAL': 'CURSO', 'LAUDO MÉDICO': 'LAUDO_MEDICO', 'OUTROS': 'OUTROS',
    'LUTO': 'LUTO',
}
TIPOS_SAUDE = {'LTS', 'RTS', 'LAUDO_MEDICO'}
NIVEIS = {
    'ENSINO FUNDAMENTAL': 1, 'ENSINO MÉDIO': 2, 'ENSINO MEDIO': 2, 'TÉCNICO': 3, 'TECNICO': 3,
    'SUPERIOR INCOMPLETO': 4, 'SUPERIOR COMPLETO': 5, 'PÓS-GRADUAÇÃO': 6, 'POS-GRADUACAO': 6,
    'ESPECIALIZAÇÃO': 6, 'MESTRADO': 7, 'DOUTORADO': 8, 'PÓS-DOUTORADO': 9,
}
PUNICOES = {'ADVERTÊNCIA': 'Advertência', 'REPREENSÃO': 'Repreensão', 'DETENÇÃO': 'Detenção', 'PRISÃO': 'Prisão'}
VAZIOS = {'', '-', '.', 'NONE', 'NÃO CONSTA', 'NAO CONSTA', '#N/A', 'N/A', 'NÃO', 'NAO'}

# ----------------------------------------------------------------------------
problemas = []   # (ome, aba, linha, matricula, nome, campo, problema, valor, acao)
def prob(r, campo, problema, valor='', acao=''):
    problemas.append((r.get('_ome', ''), r.get('_aba', ''), r.get('_linha', ''), r.get('_mat', ''),
                      r.get('_nome', ''), campo, problema, '' if valor is None else str(valor), acao))

def txt(v):
    if v is None: return None
    s = re.sub(r'\s+', ' ', str(v)).strip()
    return s or None

def vazio(v):
    return v is None or str(v).strip().upper() in VAZIOS

def sim(v):
    return v is not None and str(v).strip().upper() in ('SIM', 'S')

def norm_hdr(h):
    h = re.sub(r'\s+', ' ', str(h or '')).strip()
    h = re.sub(r'\s*:\s*$', '', h).replace('Tipo de Afastamento: (1)', 'Tipo de Afastamento (1)')
    h = h.replace('Afastamento(1)', 'Afastamento (1)')
    return 'Matrícula' if h.startswith('Matrícula') else h

def dv_ok(m):
    s = sum(int(m[i]) * (7 - i) for i in range(6))
    d = 11 - s % 11
    return (0 if d >= 10 else d) == int(m[7])

def norm_matricula(v):
    s = re.sub(r'[^0-9]', '', str(v or ''))
    if len(s) == 7: return f'{s[:6]}-{s[6]}'
    return None

def data(v):
    if isinstance(v, dt.datetime): v = v.date()
    if isinstance(v, dt.date): return v if 2000 <= v.year <= 2100 else None
    s = txt(v)
    if not s: return None
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', s)
    if m:
        try:
            d = dt.date(int(m[3]), int(m[2]), int(m[1]))
            return d if 2000 <= d.year <= 2100 else None
        except ValueError: return None
    return None

def inteiro(v):
    if v is None: return None
    if isinstance(v, (int, float)): return int(v)
    m = re.search(r'\d+', str(v))
    return int(m[0]) if m else None

def telefone(v):
    s = txt(v)
    if not s: return None, None
    d = re.sub(r'\D', '', s)
    if not d: return None, None                      # "xxx", "-", texto sem número
    if d.startswith('0') and len(d) in (11, 12): d = d[1:]   # (081) 9xxxx-xxxx
    if d.startswith('55') and len(d) in (12, 13): d = d[2:]
    if len(d) == 13 and d[:2] == d[2:4]: d = d[2:]   # DDD digitado duas vezes
    if len(d) == 11: return f'({d[:2]}) {d[2:7]}-{d[7:]}', None
    if len(d) == 10: return f'({d[:2]}) {d[2:6]}-{d[6:]}', None
    if len(d) == 9:  return f'{d[:5]}-{d[5:]}', 'telefone sem DDD'
    return s, 'telefone em formato não reconhecido'

EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')

def q(v):
    """literal SQL"""
    if v is None: return 'null'
    if isinstance(v, bool): return 'true' if v else 'false'
    if isinstance(v, int): return str(v)
    if isinstance(v, dt.date): return f"'{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"

# ----------------------------------------------------------------------------
def ler_linhas(wb):
    """Devolve lista de dicionários (um por linha de militar) de todas as abas de OME."""
    linhas = []
    # blocos da GERAL (usados para recuperar abas quebradas)
    blocos = collections.defaultdict(list)
    g = wb['GERAL'] if 'GERAL' in wb.sheetnames else None
    if g is not None:
        hdr = None
        for i, r in enumerate(g.iter_rows(values_only=True), 1):
            if r[0] == 'Posto / Graduação':
                hdr = [norm_hdr(h) for h in r]; continue
            if hdr and r[1] and r[0] != 'Atualizado em:':
                d = dict(zip(hdr, r)); blocos[txt(d.get('OME'))].append((i, d))

    for ws in wb.worksheets:
        if ws.title in ABAS_IGNORADAS: continue
        sigla = ABA_PARA_SIGLA.get(ws.title)
        if not sigla:
            print(f'  ! aba "{ws.title}" sem OME correspondente - ignorada'); continue
        hdr = [norm_hdr(c.value) for c in ws[2]]
        origem = ws.title
        if len(hdr) < 2 or hdr[1] != 'Matrícula':
            if blocos.get(sigla):
                print(f'  ! aba "{ws.title}" quebrada (#REF!) - recuperada da aba GERAL ({len(blocos[sigla])} linhas)')
                problemas.append((sigla, ws.title, '', '', '', 'Aba', 'Aba com erro #REF! na planilha',
                                  '', f'Dados recuperados da aba GERAL ({len(blocos[sigla])} militares) - conferir'))
                for ordem, (i, d) in enumerate(blocos[sigla], 1):
                    d = dict(d); d.update(_aba=ws.title + ' (via GERAL)', _ome=sigla, _linha=i, _ordem=ordem)
                    linhas.append(d)
            else:
                print(f'  ! aba "{ws.title}" sem cabeçalho reconhecível - ignorada')
            continue
        ordem = 0
        for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
            if not r or len(r) < 2 or vazio(r[1]) or str(r[1]).startswith('Matr'): continue
            d = {}
            cont = collections.Counter()
            for h, v in zip(hdr, r):
                if not h: continue
                cont[h] += 1
                d[h if cont[h] == 1 else f'{h} #{cont[h]}'] = v
            ordem += 1
            d.update(_aba=origem, _ome=sigla, _linha=i, _ordem=ordem)
            linhas.append(d)
    return linhas

def chave_ome(s):
    return re.sub(r'[\s/°º]', '', str(s or '')).upper()

# ----------------------------------------------------------------------------
def processar(linhas):
    militares, afast, saude, regs, forms, idiomas, puns = [], [], [], [], [], [], []
    secoes = collections.OrderedDict()
    duplicados = []

    for d in linhas:
        d['_nome'] = txt(d.get('Nome completo do militar'))
        m = txt(d.get('Matrícula'))
        mn = m if m and re.fullmatch(r'\d{6}-\d', m) else norm_matricula(m)
        d['_mat'] = mn or m

    # duplicidades
    por_mat = collections.defaultdict(list)
    for d in linhas: por_mat[d['_mat']].append(d)
    escolhidos = []
    for mat, grupo in por_mat.items():
        if len(grupo) == 1:
            escolhidos.append(grupo[0]); continue
        casa = [g for g in grupo if chave_ome(g.get('OME')) == chave_ome(g['_ome'])]
        fica = casa[0] if casa else grupo[0]
        criterio = 'aba coincide com a coluna OME' if casa else 'primeira ocorrência (nenhuma aba coincide com a coluna OME)'
        escolhidos.append(fica)
        duplicados.append((mat, fica['_nome'], ', '.join(f"{g['_ome']} (linha {g['_linha']})" for g in grupo),
                           txt(fica.get('OME')), fica['_ome'], criterio))
        for g in grupo:
            if g is not fica:
                prob(g, 'Matrícula', 'Militar listado em mais de uma OME',
                     f"também em {fica['_ome']}", f"Mantido no {fica['_ome']}; esta linha foi descartada")

    for d in escolhidos:
        mat = d['_mat']
        if not mat or not re.fullmatch(r'\d{6}-\d', mat):
            prob(d, 'Matrícula', 'Matrícula em formato inválido', d.get('Matrícula'), 'NÃO importado'); continue
        if not dv_ok(mat):
            prob(d, 'Matrícula', 'Dígito verificador não confere', mat, 'Importado - conferir')
        if not d['_nome'] or len(d['_nome']) < 3:
            prob(d, 'Nome', 'Nome ausente', d.get('Nome completo do militar'), 'NÃO importado'); continue

        # posto / quadro
        pg_raw = txt(d.get('Posto / Graduação')) or ''
        obs = []
        partes = pg_raw.replace('Promo. Req.', '').split()
        quadro = next((p.upper() for p in partes if p.upper() in QUADROS), None)
        base = ''.join(p for p in partes if p.upper() not in QUADROS).upper()
        posto_id = POSTOS.get(base)
        if 'Promo. Req.' in pg_raw: obs.append('Promoção a requerimento (planilha: "Subten Promo. Req.")')
        if not posto_id:
            prob(d, 'Posto/Graduação', 'Posto/graduação não reconhecido', pg_raw, 'NÃO importado'); continue

        ome = d['_ome']
        if chave_ome(d.get('OME')) != chave_ome(ome):
            prob(d, 'OME', 'Coluna OME diferente da aba onde o militar está listado', d.get('OME'),
                 f'Lotado no {ome} (aba)')

        secao = txt(d.get('Seção'))
        if vazio(secao):
            secao = None
        else:
            secoes.setdefault((ome, secao), len([k for k in secoes if k[0] == ome]) + 1)

        q_raw = (txt(d.get('Quadro')) or '').upper()
        atividade = 'OPERACIONAL' if q_raw.startswith('OPERA') else ('ADMINISTRATIVO' if q_raw.startswith('ADM') else None)
        if q_raw and not atividade and not vazio(q_raw):
            prob(d, 'Quadro', 'Atividade não reconhecida', q_raw, 'Deixado em branco')

        email = txt(d.get('Email'))
        if email:
            email = email.lower().replace(' ', '')
            if not EMAIL_RE.match(email):
                prob(d, 'E-mail', 'E-mail inválido', email, 'Deixado em branco'); email = None
        t1, e1 = telefone(d.get('Telefone - 1º Opção (DDD) 9XXXX-XXXX'))
        t2, e2 = telefone(d.get('Telefone - 2º Opção (DDD) 9XXXX-XXXX'))
        if e1: prob(d, 'Telefone 1', e1, t1, 'Importado como está')
        if e2: prob(d, 'Telefone 2', e2, t2, 'Importado como está')

        ano = inteiro(d.get('Ano Ingresso'))
        if ano is not None and not (1950 <= ano <= 2100):
            prob(d, 'Ano de ingresso', 'Ano inválido', ano, 'Deixado em branco'); ano = None

        mid = str(uuid.uuid4())
        militares.append(dict(id=mid, matricula=mat, nome=d['_nome'], posto_id=posto_id, quadro=quadro,
                              ome=ome, secao=secao, atividade=atividade, ano=ano, ordem=d['_ordem'],
                              email=email, t1=t1, t2=t2, endereco=txt(d.get('Endereço')),
                              obs='; '.join(obs) or None))

        # situações funcionais
        for flag, proc_col, tipo, det_col in [
            ('Militar Readaptado', 'Processo SEI da Readaptação', 'READAPTACAO', None),
            ('Militar com Horário Especial de Trabalho', 'Processo SEI do Horário Especial de Trabalho', 'HORARIO_ESPECIAL', None),
            ('Militar com Acumulação de Cargo Público', 'Processo SEI da Acumulação de Cargo', 'ACUMULACAO_CARGO', None),
            ('Militar PCD', None, 'PCD', 'Tipo da Deficiêrncia'),
        ]:
            proc = txt(d.get(proc_col)) if proc_col else None
            if vazio(proc): proc = None
            det = txt(d.get(det_col)) if det_col else None
            if vazio(det): det = None
            if sim(d.get(flag)):
                regs.append(dict(militar_id=mid, tipo=tipo, detalhe=det, processo=proc))
                if not proc and tipo != 'PCD':
                    prob(d, flag, 'Marcado "Sim" sem processo SEI', '', 'Importado - informar processo')
            elif proc:
                prob(d, flag, 'Processo SEI informado, mas campo marcado "Não"/vazio', proc, 'Não importado - conferir')

        # afastamentos (1) e (2)
        for n in ('1', '2'):
            tipo_raw = txt(d.get(f'Militar em Afastamento ({n})'))
            if vazio(tipo_raw): continue
            cod = TIPOS_AFAST.get(tipo_raw.upper())
            if not cod:
                prob(d, f'Afastamento ({n})', 'Tipo de afastamento não reconhecido', tipo_raw, 'Não importado'); continue
            ini_raw = d.get(f'Data de início do Afastamento ({n})')
            ini = data(ini_raw)
            dias = inteiro(d.get(f'Quantidade de dias do Afastamento ({n})'))
            fim_planilha = data(d.get(f'Término do Afastamento ({n}) *NÃO PREENCHER'))
            if ini and not dias and fim_planilha and fim_planilha >= ini:
                dias = (fim_planilha - ini).days + 1
            if not ini or not dias or dias < 1:
                prob(d, f'Afastamento ({n})', f'{tipo_raw} sem data de início ou quantidade de dias válida',
                     f'início={ini_raw} dias={d.get(f"Quantidade de dias do Afastamento ({n})")}',
                     'Não importado - cadastrar no sistema')
                continue
            if dias > 3650:
                prob(d, f'Afastamento ({n})', 'Quantidade de dias acima de 10 anos', dias, 'Não importado'); continue
            proc = txt(d.get(f'Processo SEI da Afastamento ({n})'))
            if vazio(proc): proc = None
            aid = str(uuid.uuid4())
            afast.append(dict(id=aid, militar_id=mid, tipo=cod, ini=ini, dias=dias, processo=proc,
                              obs='Curso operacional' if tipo_raw.upper() == 'CURSO OPERACIONAL' else None))
            nat_raw = (txt(d.get(f'Tipo de Afastamento ({n})')) or '').upper()
            nat = 'FISICO' if nat_raw.startswith('FÍS') or nat_raw.startswith('FIS') else ('MENTAL' if nat_raw.startswith('MENT') else None)
            cid_col = 'CID' if n == '1' else 'CID #2'
            cid = txt(d.get(cid_col))
            cid = None if vazio(cid) else cid.upper().replace(' ', '')
            if cid and not re.fullmatch(r'[A-Z][0-9]{2}(\.?[0-9A-Z]{1,2})?', cid):
                prob(d, 'CID', 'CID em formato inválido', cid, 'Descartado'); cid = None
            if cod in TIPOS_SAUDE and (nat or cid):
                saude.append(dict(afastamento_id=aid, natureza=nat, cid=cid))
            elif nat and cod not in TIPOS_SAUDE:
                pass  # "Físico" marcado em férias/dispensa etc. - descartado (sem significado)

        # punição
        pun = txt(d.get('Punição'))
        if not vazio(pun):
            nome_p = PUNICOES.get(pun.upper(), 'Outra')
            proc = txt(d.get('Processo SEI da Punição'))
            puns.append(dict(militar_id=mid, tipo=nome_p, dias=inteiro(d.get('Quantidade de dias da punição')),
                             processo=None if vazio(proc) else proc, obs=None if nome_p != 'Outra' else pun))

        # formação
        for nc, ac in (('Escolaridade', 'Área de formação'), ('Escolaridade 2', 'Área de formação 2')):
            nv = txt(d.get(nc)); area = txt(d.get(ac))
            if vazio(nv) and vazio(area): continue
            nid = NIVEIS.get((nv or '').upper())
            if not nid:
                if area and not vazio(area):
                    prob(d, nc, 'Área de formação sem nível de escolaridade', area, 'Não importado')
                elif nv:
                    prob(d, nc, 'Nível de escolaridade não reconhecido', nv, 'Não importado')
                continue
            forms.append(dict(militar_id=mid, nivel=nid, area=None if vazio(area) else area))

        # idiomas
        idi = txt(d.get('Idiomas'))
        if not vazio(idi):
            for x in re.split(r'[,;/]| e ', idi):
                x = txt(x)
                if x: idiomas.append(dict(militar_id=mid, idioma=x.capitalize()))
        # remove idiomas repetidos do mesmo militar
    vistos = set(); idiomas = [i for i in idiomas if not ((i['militar_id'], i['idioma']) in vistos or vistos.add((i['militar_id'], i['idioma'])))]
    return militares, afast, saude, regs, forms, idiomas, puns, secoes, duplicados

# ----------------------------------------------------------------------------
def values(rows, cols, fmt):
    out = []
    for i in range(0, len(rows), 500):
        out.append(',\n'.join('(' + ', '.join(fmt(r)) + ')' for r in rows[i:i + 500]))
    return out

def gerar_sql(pasta, militares, afast, saude, regs, forms, idiomas, puns, secoes):
    hoje = dt.date.today().isoformat()
    p1 = [f"""-- Gerado por importar_planilha.py em {hoje}
-- PARTE 1/2: seções, militares e registro de inclusão
begin;
select set_config('app.importacao', 'on', true);
do $$ begin
  if exists (select 1 from militar) then
    raise exception 'O banco já possui militares cadastrados. A importação inicial só pode ser feita em banco vazio.';
  end if;
end $$;
"""]
    p1.append('insert into secao (ome_id, nome, ordem)\nselect o.id, v.nome, v.ordem from (values\n' +
              ',\n'.join(f'({q(o)}, {q(n)}, {k})' for (o, n), k in secoes.items()) +
              '\n) v(ome, nome, ordem) join ome o on o.sigla = v.ome\non conflict (ome_id, nome) do nothing;\n')
    for bloco in values(militares, None, lambda m: [q(m['id']), q(m['matricula']), q(m['nome']), q(m['posto_id']),
                        q(m['quadro']), q(m['ome']), q(m['secao']), q(m['atividade']), q(m['ano']), q(m['ordem']),
                        q(m['email']), q(m['t1']), q(m['t2']), q(m['endereco']), q(m['obs'])]):
        p1.append(f"""insert into militar (id, matricula, nome, posto_id, quadro_sigla, ome_id, secao_id, atividade,
  ano_ingresso, ordem_antiguidade, email, telefone1, telefone2, endereco, observacoes)
select v.id::uuid, v.mat, v.nome, v.posto::smallint, v.quadro, o.id, s.id, v.atv::atividade_tipo, v.ano::smallint,
       v.ordem, v.email, v.t1, v.t2, v.endereco, v.obs
from (values
{bloco}
) v(id, mat, nome, posto, quadro, ome, secao, atv, ano, ordem, email, t1, t2, endereco, obs)
join ome o on o.sigla = v.ome
left join secao s on s.ome_id = o.id and s.nome = v.secao;
""")
    p1.append(f"""insert into movimentacao (militar_id, tipo, ome_destino_id, secao_destino_id, data, documento, observacao)
select id, 'INCLUSAO', ome_id, secao_id, current_date, 'Carga inicial', 'Importado da Planilha de Controle de Pessoal 2026'
from militar;
commit;
""")
    with open(os.path.join(pasta, 'seed_1_militares.sql'), 'w', encoding='utf-8') as f: f.write('\n'.join(p1))

    p2 = [f"""-- Gerado por importar_planilha.py em {hoje}
-- PARTE 2/2: afastamentos, saúde, situações funcionais, formação, idiomas e punições
begin;
select set_config('app.importacao', 'on', true);
"""]
    for bloco in values(afast, None, lambda a: [q(a['id']), q(a['militar_id']), q(a['tipo']), q(a['ini']), q(a['dias']), q(a['processo']), q(a['obs'])]):
        p2.append(f"""insert into afastamento (id, militar_id, tipo_id, data_inicio, dias, processo_sei, observacao)
select v.id::uuid, v.mil::uuid, t.id, v.ini::date, v.dias, v.proc, v.obs
from (values
{bloco}
) v(id, mil, tipo, ini, dias, proc, obs) join tipo_afastamento t on t.codigo = v.tipo;
""")
    if saude:
        p2.append('insert into afastamento_saude (afastamento_id, natureza, cid)\nselect v.id::uuid, v.nat::natureza_saude, v.cid from (values\n' +
                  ',\n'.join(f"({q(s['afastamento_id'])}, {q(s['natureza'])}, {q(s['cid'])})" for s in saude) +
                  '\n) v(id, nat, cid);\n')
    if regs:
        p2.append('insert into registro_funcional (militar_id, tipo, detalhe, processo_sei)\nselect v.mil::uuid, v.tipo::registro_funcional_tipo, v.det, v.proc from (values\n' +
                  ',\n'.join(f"({q(r['militar_id'])}, {q(r['tipo'])}, {q(r['detalhe'])}, {q(r['processo'])})" for r in regs) +
                  '\n) v(mil, tipo, det, proc);\n')
    for bloco in values(forms, None, lambda f_: [q(f_['militar_id']), q(f_['nivel']), q(f_['area'])]):
        p2.append(f"""insert into formacao (militar_id, nivel_id, area)
select v.mil::uuid, v.nivel::smallint, v.area from (values
{bloco}
) v(mil, nivel, area);
""")
    if idiomas:
        p2.append('insert into militar_idioma (militar_id, idioma)\nselect v.mil::uuid, v.idioma from (values\n' +
                  ',\n'.join(f"({q(i['militar_id'])}, {q(i['idioma'])})" for i in idiomas) +
                  '\n) v(mil, idioma) on conflict do nothing;\n')
    if puns:
        p2.append('insert into punicao (militar_id, tipo_id, dias, processo, observacao)\nselect v.mil::uuid, t.id, v.dias::int, v.proc, v.obs from (values\n' +
                  ',\n'.join(f"({q(p['militar_id'])}, {q(p['tipo'])}, {q(p['dias'])}, {q(p['processo'])}, {q(p['obs'])})" for p in puns) +
                  '\n) v(mil, tipo, dias, proc, obs) join tipo_punicao t on t.nome = v.tipo;\n')
    p2.append('commit;\n')
    with open(os.path.join(pasta, 'seed_2_detalhes.sql'), 'w', encoding='utf-8') as f: f.write('\n'.join(p2))

# ----------------------------------------------------------------------------
def gerar_relatorio(pasta, linhas, militares, afast, regs, forms, duplicados):
    wb = openpyxl.Workbook()
    cab = Font(bold=True, color='FFFFFF'); fundo = PatternFill('solid', fgColor='9B1C1C')

    def aba(ws, titulos, dados, larguras):
        ws.append(titulos)
        for c in ws[1]: c.font = cab; c.fill = fundo; c.alignment = Alignment(vertical='center', wrap_text=True)
        for r in dados: ws.append(list(r))
        for i, w in enumerate(larguras, 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

    lidos = collections.Counter(d['_ome'] for d in linhas)
    imp = collections.Counter(m['ome'] for m in militares)
    por_mil = {m['id']: m['ome'] for m in militares}
    probs = collections.Counter(p[0] for p in problemas)
    afs = collections.Counter(por_mil[a['militar_id']] for a in afast)
    sem_esc = collections.Counter(m['ome'] for m in militares if m['id'] not in {f['militar_id'] for f in forms})
    sem_tel = collections.Counter(m['ome'] for m in militares if not m['t1'])
    sem_sec = collections.Counter(m['ome'] for m in militares if not m['secao'])
    ws = wb.active; ws.title = 'Resumo por OME'
    aba(ws, ['OME', 'Linhas lidas', 'Militares importados', 'Afastamentos importados', 'Ocorrências no relatório',
             'Sem escolaridade', 'Sem telefone', 'Sem seção'],
        [(o, lidos[o], imp[o], afs[o], probs[o], sem_esc[o], sem_tel[o], sem_sec[o]) for o in sorted(lidos, key=lambda x: (-lidos[x]))] +
        [('TOTAL', sum(lidos.values()), len(militares), len(afast), len(problemas), sum(sem_esc.values()), sum(sem_tel.values()), sum(sem_sec.values()))],
        [18, 12, 12, 14, 14, 14, 12, 10])
    for c in ws[ws.max_row]: c.font = Font(bold=True)

    tipos = collections.Counter(p[6] for p in problemas)
    ws = wb.create_sheet('Tipos de ocorrência')
    aba(ws, ['Ocorrência', 'Quantidade'], tipos.most_common(), [70, 12])

    ws = wb.create_sheet('Ocorrências')
    aba(ws, ['OME', 'Aba', 'Linha', 'Matrícula', 'Nome', 'Campo', 'Ocorrência', 'Valor encontrado', 'Ação tomada'],
        sorted(problemas, key=lambda p: (str(p[0]), str(p[6]), str(p[4]))), [12, 16, 7, 11, 36, 22, 48, 30, 40])

    ws = wb.create_sheet('Duplicidades')
    aba(ws, ['Matrícula', 'Nome', 'Abas onde aparece', 'Coluna OME', 'Lotação adotada', 'Critério'],
        duplicados, [11, 38, 40, 12, 14, 50])
    wb.save(os.path.join(pasta, 'relatorio_importacao.xlsx'))

# ----------------------------------------------------------------------------
def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    arq, pasta = sys.argv[1], sys.argv[2]
    os.makedirs(pasta, exist_ok=True)
    print(f'Lendo {arq} ...')
    wb = openpyxl.load_workbook(arq, data_only=True)
    linhas = ler_linhas(wb)
    print(f'  {len(linhas)} linhas de militares lidas')
    res = processar(linhas)
    militares, afast, saude, regs, forms, idiomas, puns, secoes, duplicados = res
    gerar_sql(pasta, militares, afast, saude, regs, forms, idiomas, puns, secoes)
    gerar_relatorio(pasta, linhas, militares, afast, regs, forms, duplicados)
    print(f'  militares: {len(militares)} | seções: {len(secoes)} | afastamentos: {len(afast)} (saúde: {len(saude)})')
    print(f'  situações funcionais: {len(regs)} | formações: {len(forms)} | idiomas: {len(idiomas)} | punições: {len(puns)}')
    print(f'  duplicidades: {len(duplicados)} | ocorrências no relatório: {len(problemas)}')
    print(f'Arquivos gerados em {pasta}')

if __name__ == '__main__':
    main()
